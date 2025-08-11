import os
from typing import List
from langchain_core.documents import Document
from langchain_community.document_loaders import PyPDFLoader
from docx import Document as DocxDocument

def load_docx(file_path: str) -> List[Document]:
    try:
        doc = DocxDocument(file_path)
        full_text = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
        return [Document(page_content=full_text, metadata={"source": file_path})]
    except Exception as e:
        print(f"Error loading DOCX: {file_path} - {e}")
        return []

def load_cover_letters(folder_path: str) -> List[Document]:
    all_docs = []
    for fname in os.listdir(folder_path):
        fpath = os.path.join(folder_path, fname)
        ext = os.path.splitext(fname)[-1].lower()
        if ext == ".pdf":
            try:
                loader = PyPDFLoader(fpath)
                pdf_docs = loader.load()
                for doc in pdf_docs:
                    doc.metadata["source"] = fpath
                all_docs.extend(pdf_docs)
            except Exception as e:
                print(f"Error loading PDF: {fpath} - {e}")
        elif ext == ".docx":
            all_docs.extend(load_docx(fpath))
        else:
            print(f"Skipped unsupported file: {fpath}")
    return all_docs
